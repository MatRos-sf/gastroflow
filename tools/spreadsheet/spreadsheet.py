from datetime import datetime
from typing import Any, Optional

from openpyxl import Workbook

from tools.spreadsheet.sheets.base import BaseSheetWriter
from tools.spreadsheet.style import StyleConfig


class ManagerSpreadsheet:
    """Generates Excel spreadsheets for order reports."""

    def __init__(self, name: str, style_config: Optional[StyleConfig] = None):
        self.name = name
        self._style_config = style_config or StyleConfig()
        self._workbook = Workbook()

        if "Sheet" in self._workbook.sheetnames:
            self._workbook.remove(self._workbook["Sheet"])

    def add_sheet(
        self,
        sheet_name: str,
        SheetClass: type[BaseSheetWriter],
        data: Any,
        **kwargs: Any,
    ) -> None:
        """Create a new workbook sheet and write data using the given sheet class."""
        ws = self._workbook.create_sheet(sheet_name)
        ws.sheet_properties.defaultRowHeight = 18
        writer = SheetClass(ws, style_config=self._style_config)
        writer.write_all(data, **kwargs)

    def save(self) -> None:
        """Save the workbook to file."""
        import os

        directory = os.path.dirname(self.name)
        if directory and not os.path.exists(directory):
            os.makedirs(directory)
        self._workbook.save(self.name)

    @staticmethod
    def create_report_name(
        from_date: Optional[datetime] = None,
        to_date: Optional[datetime] = None,
    ) -> str:
        """Generate a report filename based on date range."""
        if from_date and to_date and to_date.date() != from_date.date():
            return f"raport_{from_date.date()}_{to_date.date()}.xlsx"
        if from_date:
            return f"raport_{from_date.date()}.xlsx"
        if to_date:
            return f"raport_do_{to_date.date()}.xlsx"
        return "raport_calosciowy.xlsx"
