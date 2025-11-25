__all__ = ["TableWriterSheet"]
from decimal import Decimal
from typing import Optional

from tools.spreadsheet.style import CellStyle, StyleConfig


class TableWriterSheet:
    """Handles writing tabular data to an Excel sheet."""

    COLUMN_HEADERS = ("Name", "Quantity", "Revenue")

    def __init__(
        self,
        sheet,
        start_row: int = 1,
        start_col: int = 1,
        style_config: Optional[StyleConfig] = None,
    ):
        self._sheet = sheet
        self._start_row = start_row
        self._start_col = start_col
        self._current_row = start_row
        self._current_col = start_col
        self._max_name_length = 0
        self._style_config = style_config or StyleConfig()

    def _apply_style(self, cell, style: CellStyle) -> None:
        """Apply a CellStyle to a cell."""
        if style.font:
            cell.font = style.font
        if style.fill:
            cell.fill = style.fill
        if style.alignment:
            cell.alignment = style.alignment
        if style.border:
            cell.border = style.border
        if style.number_format:
            cell.number_format = style.number_format

    def write_headers(
        self,
        main_title: str,
        title_style: Optional[CellStyle] = None,
        header_style: Optional[CellStyle] = None,
    ) -> None:
        """Write column headers to the sheet."""
        # Use provided styles or defaults
        title_style = title_style or self._style_config.get_title_style()
        header_style = header_style or self._style_config.get_header_style()

        # Write main title (merged across 3 columns)
        cell = self._sheet.cell(row=self._current_row, column=self._current_col)
        cell.value = main_title
        self._apply_style(cell, title_style)

        # Merge cells for title
        self._sheet.merge_cells(
            start_row=self._current_row,
            start_column=self._current_col,
            end_row=self._current_row,
            end_column=self._current_col + 2,
        )
        self._current_row += 1

        # Write column headers
        for idx, header in enumerate(self.COLUMN_HEADERS):
            cell = self._sheet.cell(
                row=self._current_row, column=self._current_col + idx
            )
            cell.value = header
            self._apply_style(cell, header_style)
        self._current_row += 1

    def write_row(
        self,
        name: str,
        quantity: int,
        revenue: Decimal,
        name_style: Optional[CellStyle] = None,
        quantity_style: Optional[CellStyle] = None,
        revenue_style: Optional[CellStyle] = None,
    ) -> None:
        """Write a single data row."""
        # Use provided styles or defaults
        name_style = name_style or self._style_config.get_data_style("left")
        quantity_style = quantity_style or self._style_config.get_data_style("center")
        revenue_style = revenue_style or self._style_config.get_data_style("right")

        # Add number format to revenue
        if not revenue_style.number_format:
            revenue_style.number_format = self._style_config.FORMAT_DECIMAL

        # Name cell
        name_cell = self._sheet.cell(
            row=self._current_row, column=self._current_col, value=name
        )
        self._apply_style(name_cell, name_style)

        # Quantity cell
        quantity_cell = self._sheet.cell(
            row=self._current_row, column=self._current_col + 1, value=quantity
        )
        self._apply_style(quantity_cell, quantity_style)

        # Revenue cell
        revenue_cell = self._sheet.cell(
            row=self._current_row, column=self._current_col + 2, value=float(revenue)
        )
        self._apply_style(revenue_cell, revenue_style)

        self._current_row += 1
        self._max_name_length = max(self._max_name_length, len(name))

    def write_summary(
        self,
        total_quantity: int,
        total_revenue: Decimal,
        label: str = "Summary",
        label_style: Optional[CellStyle] = None,
        quantity_style: Optional[CellStyle] = None,
        revenue_style: Optional[CellStyle] = None,
    ) -> None:
        """Write summary row with styling."""
        # Use provided styles or defaults
        label_style = label_style or self._style_config.get_summary_style("center")
        quantity_style = quantity_style or self._style_config.get_summary_style(
            "center"
        )
        revenue_style = revenue_style or self._style_config.get_summary_style("right")

        # Add number format to revenue
        if not revenue_style.number_format:
            revenue_style.number_format = self._style_config.FORMAT_DECIMAL

        # Label cell
        summary_cell = self._sheet.cell(
            row=self._current_row, column=self._current_col, value=label
        )
        self._apply_style(summary_cell, label_style)

        # Quantity cell
        quantity_cell = self._sheet.cell(
            row=self._current_row, column=self._current_col + 1, value=total_quantity
        )
        self._apply_style(quantity_cell, quantity_style)

        # Revenue cell
        revenue_cell = self._sheet.cell(
            row=self._current_row,
            column=self._current_col + 2,
            value=float(total_revenue),
        )
        self._apply_style(revenue_cell, revenue_style)

        self._current_row += 1

    def adjust_columns(
        self,
        name_width: Optional[int] = None,
        quantity_width: int = 12,
        revenue_width: int = 15,
    ) -> None:
        """Adjust column widths based on content."""
        from openpyxl.utils import get_column_letter

        # Name column - based on content or provided width
        col_letter = get_column_letter(self._current_col)
        width = name_width or max(self._max_name_length + 2, 20)
        self._sheet.column_dimensions[col_letter].width = width

        # Quantity column
        col_letter = get_column_letter(self._current_col + 1)
        self._sheet.column_dimensions[col_letter].width = quantity_width

        # Revenue column
        col_letter = get_column_letter(self._current_col + 2)
        self._sheet.column_dimensions[col_letter].width = revenue_width

    def start_new_table(self, column_gap: int = 2) -> None:
        """Start a new table with specified offset from current position."""
        self._current_row = self._start_row
        self._current_col += len(self.COLUMN_HEADERS) + column_gap
        self._max_name_length = 0
