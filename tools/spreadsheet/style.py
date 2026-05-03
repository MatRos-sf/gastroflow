from dataclasses import dataclass
from typing import Optional

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@dataclass
class CellStyle:
    """Configuration for cell styling."""

    font: Optional[Font] = None
    fill: Optional[PatternFill] = None
    alignment: Optional[Alignment] = None
    border: Optional[Border] = None
    number_format: Optional[str] = None


class StyleConfig:
    """Centralized style configuration for the spreadsheet."""

    # Colors
    COLOR_TITLE_BG = "4472C4"
    COLOR_HEADER_BG = "5B9BD5"
    COLOR_SUMMARY_BG = "FFEB9C"
    COLOR_WHITE = "FFFFFF"

    # Fonts
    FONT_TITLE = Font(bold=True, size=14, color=COLOR_WHITE)
    FONT_HEADER = Font(bold=True, size=11, color=COLOR_WHITE)
    FONT_BOLD = Font(bold=True, size=11)
    FONT_NORMAL = Font(size=11)

    # Fills
    FILL_TITLE = PatternFill(
        start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type="solid"
    )
    FILL_HEADER = PatternFill(
        start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid"
    )
    FILL_SUMMARY = PatternFill(
        start_color=COLOR_SUMMARY_BG, end_color=COLOR_SUMMARY_BG, fill_type="solid"
    )

    # Alignments
    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

    # Borders
    BORDER_THIN = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    BORDER_MEDIUM = Border(
        left=Side(style="medium"),
        right=Side(style="medium"),
        top=Side(style="medium"),
        bottom=Side(style="medium"),
    )

    # Number formats
    FORMAT_DECIMAL = "#,##0.00"
    FORMAT_INTEGER = "#,##0"

    # Default styles
    @classmethod
    def get_title_style(cls) -> CellStyle:
        return CellStyle(
            font=cls.FONT_TITLE, fill=cls.FILL_TITLE, alignment=cls.ALIGN_CENTER
        )

    @classmethod
    def get_header_style(cls) -> CellStyle:
        return CellStyle(
            font=cls.FONT_HEADER,
            fill=cls.FILL_HEADER,
            alignment=cls.ALIGN_CENTER,
            border=cls.BORDER_THIN,
        )

    @classmethod
    def get_data_style(cls, alignment: str = "left") -> CellStyle:
        align_map = {
            "left": cls.ALIGN_LEFT,
            "center": cls.ALIGN_CENTER,
            "right": cls.ALIGN_RIGHT,
        }
        return CellStyle(
            font=cls.FONT_NORMAL,
            alignment=align_map.get(alignment, cls.ALIGN_LEFT),
            border=cls.BORDER_THIN,
        )

    @classmethod
    def get_summary_style(cls, alignment: str = "center") -> CellStyle:
        align_map = {
            "left": cls.ALIGN_LEFT,
            "center": cls.ALIGN_CENTER,
            "right": cls.ALIGN_RIGHT,
        }
        return CellStyle(
            font=cls.FONT_BOLD,
            fill=cls.FILL_SUMMARY,
            alignment=align_map.get(alignment, cls.ALIGN_CENTER),
            border=cls.BORDER_MEDIUM,
        )
